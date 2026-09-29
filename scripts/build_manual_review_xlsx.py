import csv
import json
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = "/workspace/missed-abortion"
CSV_PATH = os.path.join(ROOT, "output", "random_80_manual_review.csv")
OUT = os.path.join(ROOT, "output", "random_80_manual_review.xlsx")

FIELD_MEANINGS = {
    "G": "הריונות — סה\"כ הריונות, כולל ההריון הנוכחי",
    "P": "לידות מעבר לכשירות חיים (כולל ניתוחים קיסריים, לא כולל הפלות)",
    "NVD": "לידות רגילות (וגינליות) בלבד",
    "CS": "ניתוחים קיסריים",
    "Ab_total": "סה\"כ הפלות בהיסטוריה (ספונטניות + יזומות + הפלה נדחית וכו')",
    "prev_D_C": "מספר גרידות מתועדות בהיסטוריה",
    "IVF": "0=הריון טבעי, 1=IVF רגיל, 2=IVF עם PGD (בדיקה גנטית טרום השרשה)",
    "smoking": "מעשנת (כן/לא)",
    "Rh": "סוג דם מלא (כולל +/-)",
    "HGB": "המוגלובין (ספירת דם)",
    "PLT": "טסיות דם",
    "WBC": "ספירת תאי דם לבנים (לויקוציטים)",
    "hcg_value": "ערך BHCG כמותי (מספרי)",
    "hcg_qualitative": "תוצאת BHCG איכותית (חיובי/שלילי) — כשלא נבדק ערך מספרי",
    "GA_LMP_weeks": "שבוע הריון לפי תאריך וסת אחרונה (LMP)",
    "GA_US_weeks": "שבוע הריון לפי מדידת סונר (CRL)",
    "US_finding": "ממצא סונר בסולם 1-4: 1=שק הריון בלבד, 2=שק חלמון בלי קוטב עוברי, 3=קוטב עוברי בלי דופק שתועד, 4=קוטב עוברי עם דופק שתועד בעבר ונעלם",
    "medical_treatment": "ניתן טיפול תרופתי (מייפיג'ין/ציטוטק)",
    "D_C_treatment": "גרידה בוצעה בפועל",
    "elective_D_C": "נקבעה גרידה מתוכננת (לא דחופה)",
    "emergency_D_C": "בוצעה גרידה דחופה (עקב דימום/חוסר יציבות)",
    "hysteroscopy": "בוצעה היסטרוסקופיה",
    "FU1_gestational": "ביקורת מעקב ראשונה: שק הריון עדיין נצפה",
    "FU1_lining_over15mm": "ביקורת מעקב ראשונה: עובי רירית הרחם מעל 15 מ\"מ",
    "FU1_bleeding": "ביקורת מעקב ראשונה: דימום",
    "FU1_pain": "ביקורת מעקב ראשונה: כאב",
    "FU2_gestational": "ביקורת מעקב שנייה: שק הריון עדיין נצפה",
    "FU2_lining_over15mm": "ביקורת מעקב שנייה: עובי רירית הרחם מעל 15 מ\"מ",
    "FU2_bleeding": "ביקורת מעקב שנייה: דימום",
    "FU2_pain": "ביקורת מעקב שנייה: כאב",
    "delayed_decision_followup": "המטופלת ביקשה לדחות את ההחלטה על הטיפול (לחשוב ולחזור)",
}

GENERAL_MEANINGS = [
    ("טקסט_קליני_מקורי_מלא", "הטקסט הגולמי המלא של הרשומות הקליניות שקודדו עבור אפיזודה זו (אישה אחת, אירוע אחד) — עשוי לכלול כמה רישומים כרונולוגיים."),
    ("episode_id", "מזהה ייחודי של האפיזודה במערכת (RNDEP=קבוצת דגימה ראשונה, RND2EP=קבוצת דגימה שנייה)."),
    ("admission_nos", "מספר/י האשפוז המקוריים ששייכים לאפיזודה — יותר ממספר אחד אם כמה אשפוזים נקשרו יחד כאותו אירוע."),
    ("age", "גיל האישה."),
    ("entry_date", "תאריך ושעת תחילת האפיזודה (הרישום הראשון)."),
    ("discharge_date", "תאריך ושעת סיום האפיזודה (הרישום האחרון הידוע)."),
    ("רלוונטי_למחקר", "\"לא\" כאשר האישה בחרה בגרידה מלכתחילה בלי לנסות טיפול תרופתי, או שהאפיזודה הוחרגה כי היא בכלל לא הפלה נדחית אמיתית. אחרת \"כן\"."),
    ("סיבת_אי_רלוונטיות", "הסבר קצר למה סומן \"לא\" בעמודת הרלוונטיות."),
    ("עדיפות_בדיקה", "דירוג לתעדוף בדיקה רפואית: גבוהה (חשד לערבוב רשומות, או שהקודדים לא הסכימו על הסיווג הראשי), בינונית (3+ אי-הסכמות בשדות משניים), רגילה (שאר המקרים)."),
    ("סיבת_עדיפות", "מה בדיוק גרם לדירוג העדיפות."),
    ("הכרעה_סוג_תוצאה", "הסיווג הסופי של האפיזודה, שנקבע לפי רוב קולות מבין 3 קודדי ה-AI: הצלחה / כישלון / פתוח / לא רלוונטי / הוחרג."),
    ("הכרעה_סיבת_החרגה", "אם הסיווג הוא \"הוחרג\" — הסיבה הספציפית (למשל הריון חוץ רחמי, תאומים, הפסקת הריון יזומה וכו')."),
    ("הכרעה_ציטוט_תומך", "ציטוט קצר מתוך הטקסט המקורי שתומך בסיווג שנבחר."),
    ("קודד1_סוג_תוצאה / קודד2_סוג_תוצאה / קודד3_סוג_תוצאה", "מה כל אחד משלושת קודדי ה-AI קבע בנפרד, לפני ההכרעה הסופית (רוב קולות)."),
    ("הסכמה_מלאה_בסיווג", "האם שלושת הקודדים הסכימו ביניהם על הסיווג הראשי (\"כן\") או שהיה פיצול (\"לא\")."),
    ("חשד_לערבוב_רשומות", "האם קודד/ים זיהו חשד שרישום ישן/לא-קשור נכלל בטעות בתוך האפיזודה הזו."),
    ("הערת_ערבוב_רשומות", "פירוט החשד לערבוב רשומות, אם סומן."),
    ("הערת_קודד", "הערה חופשית שקודד כתב על נקודה גבולית/לא חד-משמעית שדורשת בדיקה אנושית."),
    ("[שם שדה]__הכרעה", "הערך הסופי שנקבע לשדה הקליני הזה (רוב קולות מבין 3 הקודדים)."),
    ("[שם שדה]__קודד1_2_3", "מה כל אחד משלושת הקודדים ענה בנפרד לשדה זה, בפורמט \"קודד1 / קודד2 / קודד3\"."),
    ("מספר_אי_הסכמות", "כמה שדות (מתוך כלל השדות הקליניים) לא הייתה עליהם הסכמה מלאה בין שלושת הקודדים."),
    ("שדות_עם_אי_הסכמה", "רשימת שמות השדות הספציפיים שיש בהם אי-הסכמה בין הקודדים."),
    ("האם_נכון_(למילוי_ע\"י_רופא/ה)", "עמודה ריקה למילוי ידני: לסמן אם הקידוד (בכללותו, או של שדה ספציפי שבודקים) נכון."),
    ("אם_לא_נכון__הערך_הנכון", "עמודה ריקה למילוי ידני: אם הקידוד שגוי, מה הערך הנכון."),
    ("הערת_רופא/ה", "עמודה ריקה להערות חופשיות של הרופא/ה הבודק/ת."),
]

# ---------------------------------------------------------------------------
rows = list(csv.reader(open(CSV_PATH, encoding="utf-8-sig")))
header, data_rows = rows[0], rows[1:]

wb = openpyxl.Workbook()
FONT = "Arial"
header_font = Font(name=FONT, bold=True, size=10)
cell_font = Font(name=FONT, size=10)
legend_font = Font(name=FONT, size=10)
legend_bold = Font(name=FONT, size=10, bold=True)

# --- Sheet 1: legend / column explanations (first, so it's the first thing opened) ---
ws_legend = wb.active
ws_legend.title = "הסבר עמודות"
ws_legend.sheet_view.rightToLeft = True
ws_legend.freeze_panes = "A2"
ws_legend.cell(row=1, column=1, value="עמודה").font = header_font
ws_legend.cell(row=1, column=2, value="הסבר").font = header_font
ws_legend.column_dimensions["A"].width = 34
ws_legend.column_dimensions["B"].width = 95
r = 2
for name, meaning in GENERAL_MEANINGS:
    ws_legend.cell(row=r, column=1, value=name).font = legend_bold
    c = ws_legend.cell(row=r, column=2, value=meaning)
    c.font = legend_font
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
    ws_legend.cell(row=r, column=1).alignment = Alignment(vertical="top", horizontal="right", readingOrder=2)
    r += 1

r += 1
ws_legend.cell(row=r, column=1, value="שמות השדות הקליניים (__הכרעה / __קודד1_2_3)").font = header_font
r += 1
for field, meaning in FIELD_MEANINGS.items():
    ws_legend.cell(row=r, column=1, value=field).font = legend_bold
    c = ws_legend.cell(row=r, column=2, value=meaning)
    c.font = legend_font
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
    ws_legend.cell(row=r, column=1).alignment = Alignment(vertical="top", horizontal="right", readingOrder=2)
    r += 1

# --- Sheet 2: data ---
ws = wb.create_sheet("נתונים")
ws.sheet_view.rightToLeft = True
ws.freeze_panes = "B2"
for j, h in enumerate(header, start=1):
    ws.cell(row=1, column=j, value=h).font = header_font

priority_fill = {
    "גבוהה": PatternFill(start_color="F6DEDE", end_color="F6DEDE", fill_type="solid"),
    "בינונית": PatternFill(start_color="FCEFD1", end_color="FCEFD1", fill_type="solid"),
}
priority_col = header.index("עדיפות_בדיקה")

for i, row in enumerate(data_rows, start=2):
    fill = priority_fill.get(row[priority_col])
    for j, val in enumerate(row, start=1):
        cell = ws.cell(row=i, column=j, value=val)
        cell.font = cell_font
        cell.alignment = Alignment(wrap_text=False, vertical="top", horizontal="right", readingOrder=2)
        if fill is not None:
            cell.fill = fill

col_widths = {"A": 90, "B": 12, "C": 16, "D": 6, "E": 20, "F": 20, "G": 10, "H": 26, "I": 10, "J": 26}
for letter, w in col_widths.items():
    ws.column_dimensions[letter].width = w
for j in range(11, len(header) + 1):
    letter = get_column_letter(j)
    ws.column_dimensions[letter].width = 16

wb.save(OUT)
print("saved", OUT)
